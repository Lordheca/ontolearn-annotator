"""Build the class list and the per-image label CSV for the EPP dataset.

Reads the two EPP spreadsheets and writes:

* ``classes.csv``  (``code;name_en;name_ja``) - the class list, in the
  hierarchical order of the legend. It is the single source of truth for the
  classes: ``epp_dataset_builder.py`` reads it, and the training notebook saves
  it next to ``model.keras`` so ``playground.py`` uses the same order.
* ``data/5k_epp_dataset.csv`` (``image_name;label``) - one row per image,
  ``label`` being a class code from ``classes.csv``.

A class is any legend code with at least one labelled image. Rows whose label
is not in the legend (e.g. ``X``) are skipped and reported.

Usage (from examples/water_crystal_classification):

    pip install openpyxl
    python datasets/epp_dataset/prepare_labels.py \
        --legend path/to/EPP_Classification_202308_V1.xlsx \
        --labels path/to/5k_label_2023August_update_v12.xlsx
"""

import argparse
import csv
import re
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent


def normalize_code(value) -> str:
    """Turn a raw spreadsheet code into a canonical string.

    Handles the inconsistencies found in both files: floats (``1.1``), ints
    (``3``), stray spaces (``"1.7 .5"``), trailing dots (``"2.0. "``) and a
    ``.0`` suffix on top-level codes (``"3.0"`` -> ``"3"``).
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    code = re.sub(r"\s+", "", str(value)).strip(".")
    if re.fullmatch(r"\d+\.0", code):
        code = code[:-2]
    return code


def read_legend(path: Path):
    """Return ``[(code, name_en, name_ja), ...]`` in legend order."""
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    legend = []
    for code, name_ja, name_en in sheet.iter_rows(min_col=1, max_col=3, values_only=True):
        code = normalize_code(code)
        if not code:
            continue
        legend.append((code, (name_en or "").strip(), (name_ja or "").strip()))
    return legend


def read_labels(path: Path):
    """Return ``[(image_name, code), ...]`` from the ``Img_name``/``Label`` columns."""
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    name_col, label_col = header.index("Img_name"), header.index("Label")
    labels = []
    for row in rows:
        name = row[name_col]
        if not name:
            continue
        labels.append((str(name).strip(), normalize_code(row[label_col])))
    return labels


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--legend", required=True, type=Path)
    parser.add_argument("--labels", required=True, type=Path)
    parser.add_argument("--classes-out", type=Path, default=HERE / "classes.csv")
    parser.add_argument("--dataset-out", type=Path, default=HERE / "data" / "5k_epp_dataset.csv")
    args = parser.parse_args()

    legend = read_legend(args.legend)
    labels = read_labels(args.labels)

    duplicates = [name for name, n in Counter(name for name, _ in labels).items() if n > 1]
    if duplicates:
        raise SystemExit(f"Duplicate image names in {args.labels}: {duplicates[:10]}")

    counts = Counter(code for _, code in labels)
    known = {code for code, _, _ in legend}
    classes = [entry for entry in legend if counts[entry[0]] > 0]
    kept = [(name, code) for name, code in labels if code in known]
    skipped = Counter(code for _, code in labels if code not in known)

    args.classes_out.parent.mkdir(parents=True, exist_ok=True)
    with open(args.classes_out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["code", "name_en", "name_ja"])
        writer.writerows(classes)

    args.dataset_out.parent.mkdir(parents=True, exist_ok=True)
    with open(args.dataset_out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["image_name", "label"])
        writer.writerows(kept)

    print(f"{len(classes)} classes -> {args.classes_out}")
    for code, name_en, _ in classes:
        print(f"  {code:<6} {counts[code]:>5}  {name_en}")
    unused = [code for code, _, _ in legend if counts[code] == 0]
    print(f"Legend codes with no images (group headers, not classes): {unused}")
    print(f"{len(kept)} labelled images -> {args.dataset_out}")
    if skipped:
        print(f"Skipped, label not in legend: {dict(skipped)}")


if __name__ == "__main__":
    main()
